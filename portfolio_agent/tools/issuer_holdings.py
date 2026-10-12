"""
Issuer holdings files — the full, daily holdings list an ETF issuer publishes for its own fund (iShares CSV, SPDR/SSGA spreadsheet).

Why: yfinance gives a fund's top 10 holdings only, and N-PORT lags a quarter. An issuer file is current (daily), complete, and carries
each holding's TICKER and SECTOR — so it needs no identifier mapping and gives a real sector mix.

Which funds: rows of fund_issuer_sources (one per fund ticker: issuer, url, enabled). The table is SEEDED DISABLED and admin-edited; this
module downloads nothing for a fund unless its row exists AND is enabled. Only https URLs on the issuer's own domain are fetched.

Layouts handled (both are "some preamble, a header row, the holdings, then footer text"; the header row is FOUND, not assumed):
  iShares   header  Ticker, Name, Sector, Asset Class, Market Value, Weight (%), ...   preamble row "Fund Holdings as of","Oct 10, 2026"
  SPDR      header  Name, Identifier, SEDOL, Weight, Sector, Shares Held, Local Currency   preamble row "Holdings:","As of 09-Oct-2026"
Weights are percents in both and become fractions here.

What becomes what
  * a security row (equity / fixed income / ETF-of-something) -> a holding (symbol = the file's ticker when it is one, else None);
  * cash and money-market rows -> the asset-class mix only; futures / swaps / negative weights -> derivative_short_weight only;
  * sector_weights = each sector's share of the EQUITY part (the convention the look-through engine uses), with GICS names translated to the
    names the rest of the app uses (Information Technology -> Technology ...).
No network in tests: the HTTP session is injectable and the fixtures in tests/fixtures/issuer are SYNTHETIC (written from the layouts above,
not downloaded), so a real file may differ in details this parser does not yet know about; an unreadable file is a reason code, never an exception.
"""
from __future__ import annotations

import csv
import io
import os
import re
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from urllib.parse import urlparse

from portfolio_agent.tools import fund_data_db as fdb
from portfolio_agent.tools.fund_data_providers import (
    RC_FUND_ASSET_CLASSES_MISSING, RC_FUND_FETCH_FAILED, RC_FUND_HOLDINGS_MISSING, RC_ISSUER_NOT_CONFIGURED,
    FundData, FundDataProvider, FundHolding, FundProfile,
)

ISSUER_HOSTS = {"ishares": ("ishares.com",), "spdr": ("ssga.com",)}
MAX_BYTES = 25 * 1024 * 1024
_SPACING_SECONDS = 1.0                                            # one issuer request a second, process-wide

_GICS_TO_APP = {
    "information technology": "Technology", "technology": "Technology", "health care": "Healthcare", "healthcare": "Healthcare",
    "financials": "Financial Services", "financial": "Financial Services", "financial services": "Financial Services",
    "consumer discretionary": "Consumer Cyclical", "consumer cyclical": "Consumer Cyclical",
    "communication": "Communication Services", "communication services": "Communication Services",
    "communications": "Communication Services", "consumer staples": "Consumer Defensive", "consumer defensive": "Consumer Defensive",
    "materials": "Basic Materials", "basic materials": "Basic Materials", "energy": "Energy", "industrials": "Industrials",
    "utilities": "Utilities", "real estate": "Real Estate",
}


class IssuerFileError(Exception):
    """The file could not be fetched or understood."""


def validate_issuer_url(issuer: str, url: str) -> str:
    """https only, and the host must be the issuer's own domain (or a subdomain of it). Returns the url; raises ValueError."""
    u = urlparse((url or "").strip())
    hosts = ISSUER_HOSTS.get(issuer)
    if hosts is None:
        raise ValueError(f"unknown issuer {issuer!r}; known: {sorted(ISSUER_HOSTS)}")
    if u.scheme != "https" or not u.hostname:
        raise ValueError("the URL must be https")
    h = u.hostname.lower()
    if not any(h == d or h.endswith("." + d) for d in hosts):
        raise ValueError(f"the URL host must be on {' / '.join(hosts)} for {issuer}")
    return url.strip()


# ── Reading the file into rows ────────────────────────────────────────────────

def _rows_from_bytes(raw: bytes) -> list[list[str]]:
    """CSV (any text encoding with an optional BOM) or .xlsx (a zip: starts with PK) -> a list of rows of strings."""
    if raw[:2] == b"PK":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            return [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
        except Exception as exc:
            raise IssuerFileError(f"unreadable spreadsheet: {type(exc).__name__}") from None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def _num(v: str) -> float | None:
    v = (v or "").replace(",", "").replace("$", "").replace("%", "").strip()
    if v in ("", "-", "--", "N/A", "n/a"):
        return None
    try:
        return float(v)
    except ValueError:
        return None


_DATE_FORMATS = ("%b %d, %Y", "%B %d, %Y", "%d-%b-%Y", "%m/%d/%Y", "%Y-%m-%d", "%d %b %Y")


def _parse_date(text: str) -> str | None:
    t = re.sub(r"(?i)^as of\s*", "", (text or "").strip()).strip().strip('"')
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            continue
    return None


@dataclass
class IssuerRow:
    ticker: str | None
    name: str | None
    sector: str | None
    asset_class: str | None
    weight: float                    # fraction


@dataclass
class IssuerFile:
    as_of: str | None
    rows: list[IssuerRow]


def parse_issuer_file(raw: bytes) -> IssuerFile:
    """Find the header row (a row that has a Name column and a Weight column), read holdings until the table ends, find the as-of date in
    the preamble. Raises IssuerFileError when there is no recognisable table."""
    table = _rows_from_bytes(raw)
    as_of = None
    head_i, cols = None, {}
    for i, row in enumerate(table[:60]):
        low = [c.lower() for c in row]
        for j, c in enumerate(row):
            if as_of is None and ("holdings as of" in c.lower() or c.lower() in ("holdings:", "holdings as of:", "as of")):
                as_of = next((_parse_date(x) for x in row[j + 1:j + 3] if _parse_date(x)), None)
            elif as_of is None and re.match(r"(?i)^holdings:?\s*as of", c):
                as_of = _parse_date(c.split("as of", 1)[-1] if "as of" in c else c)
        w = next((k for k, c in enumerate(low) if c in ("weight (%)", "weight", "% of net assets", "weight(%)")), None)
        n = next((k for k, c in enumerate(low) if c == "name"), None)
        if w is not None and n is not None:
            head_i = i
            cols = {"weight": w, "name": n,
                    "ticker": next((k for k, c in enumerate(low) if c in ("ticker", "identifier")), None),
                    "sector": next((k for k, c in enumerate(low) if c == "sector"), None),
                    "asset_class": next((k for k, c in enumerate(low) if c == "asset class"), None)}
            break
    if head_i is None:
        raise IssuerFileError("no holdings table (a Name column and a Weight column) found")
    get = lambda row, k: (row[k] if k is not None and k < len(row) else "").strip()
    rows: list[IssuerRow] = []
    for row in table[head_i + 1:]:
        if not any(row):
            if rows:
                break                                              # the blank line after the table; footer text follows
            continue
        w = _num(get(row, cols["weight"]))
        if w is None:
            continue
        rows.append(IssuerRow(ticker=get(row, cols["ticker"]) or None, name=get(row, cols["name"]) or None,
                              sector=get(row, cols["sector"]) or None, asset_class=get(row, cols["asset_class"]) or None,
                              weight=w / 100.0))
    if not rows:
        raise IssuerFileError("the holdings table is empty")
    return IssuerFile(as_of=as_of, rows=rows)


# ── Rows -> FundData ──────────────────────────────────────────────────────────

_CASH = re.compile(r"\b(cash|money market|usd|us dollar|u\.s\. dollar|collateral|margin)\b", re.I)
_DERIV = re.compile(r"\b(future|futures|swap|swaps|option|options|forward|fx contract|derivative)\b", re.I)


def _has_sector(r: IssuerRow) -> bool:
    """A real equity sector (so "Unassigned" / "Cash and/or Derivatives" do not make a cash line look like a stock)."""
    return (r.sector or "").strip().lower() in _GICS_TO_APP


def _kind(r: IssuerRow) -> str:
    """security | cash | derivative — from the file's asset class when it has one, else from the name."""
    ac = (r.asset_class or "").lower()
    if r.weight < 0 or "future" in ac or "swap" in ac or "derivative" in ac or "option" in ac:
        return "derivative"
    if "cash" in ac or "money market" in ac:
        return "cash"
    if _DERIV.search(r.name or ""):
        return "derivative"
    if _CASH.search(r.name or "") and not _has_sector(r):
        return "cash"
    return "security"


def _class_of(r: IssuerRow) -> str:
    ac = (r.asset_class or "").lower()
    if "fixed" in ac or "bond" in ac:
        return "bond"
    if "equity" in ac or "preferred" in ac or (not ac and _has_sector(r)):
        return "stock"
    return "other"


# Class-share tickers the way issuers print them -> the way market data (yfinance) writes them. A small, explicit table plus "." -> "-":
# nothing is inferred beyond it, so a ticker not in it stays exactly as the file wrote it.
_CLASS_SHARE_TICKERS = {"BRKB": "BRK-B", "BRKA": "BRK-A", "BFB": "BF-B", "BFA": "BF-A", "LENB": "LEN-B"}


def normalize_issuer_ticker(sym: str | None) -> str | None:
    t = (sym or "").strip().upper()
    if t in ("", "-", "N/A", "--") or not re.match(r"^[A-Z0-9.\-]{1,10}$", t):
        return None
    t = t.replace(".", "-")
    return _CLASS_SHARE_TICKERS.get(t, t)


def issuer_file_to_fund_data(ticker: str, f: IssuerFile) -> FundData:
    data = FundData(ticker=ticker.upper(), source="issuer_csv", as_of=f.as_of)
    holdings: list[FundHolding] = []
    classes: dict[str, float] = {}
    sectors: dict[str, float] = {}
    deriv = 0.0
    for r in f.rows:
        k = _kind(r)
        if k == "derivative":
            deriv += r.weight
            continue
        if k == "cash":
            classes["cash"] = classes.get("cash", 0.0) + r.weight
            continue
        if r.weight <= 0:
            continue
        sym = normalize_issuer_ticker(r.ticker)
        holdings.append(FundHolding(symbol=sym, name=r.name, weight=round(r.weight, 12)))
        c = _class_of(r)
        classes[c] = classes.get(c, 0.0) + r.weight
        if c == "stock" and r.sector:
            name = _GICS_TO_APP.get(r.sector.strip().lower())
            if name:
                sectors[name] = sectors.get(name, 0.0) + r.weight
    holdings.sort(key=lambda h: -h.weight)
    data.holdings = holdings
    data.holdings_coverage = round(min(1.0, sum(h.weight for h in holdings)), 12)
    data.mapped_coverage = round(min(1.0, sum(h.weight for h in holdings if h.symbol)), 12)
    data.derivative_short_weight = round(deriv, 12)
    data.asset_class_weights = {k: round(v, 12) for k, v in classes.items() if v > 0}
    eq = sum(sectors.values())
    data.sector_weights = {k: round(v / eq, 12) for k, v in sectors.items() if v > 0} if eq > 0 else {}
    data.profile = FundProfile()
    if not holdings:
        data.reason_codes.append(RC_FUND_HOLDINGS_MISSING)
    if not data.asset_class_weights:
        data.reason_codes.append(RC_FUND_ASSET_CLASSES_MISSING)
    return data


# ── Provider ──────────────────────────────────────────────────────────────────

class _Spacing:
    def __init__(self, interval, *, clock=time.monotonic, sleep=time.sleep):
        self._i, self._clock, self._sleep, self._next, self._lock = interval, clock, sleep, 0.0, threading.Lock()

    def wait(self):
        with self._lock:
            now = self._clock()
            start = max(now, self._next)
            self._next = start + self._i
        if start > now:
            self._sleep(start - now)


_shared_spacing = _Spacing(_SPACING_SECONDS)


class IssuerCSVProvider(FundDataProvider):
    """Holdings straight from the issuer's own file, for funds with an ENABLED fund_issuer_sources row. Never raises."""
    name = "issuer_csv"

    def __init__(self, *, session=None, spacing=None, timeout: float = 30.0):
        self._session, self._spacing, self._timeout = session, spacing or _shared_spacing, timeout

    def _download(self, url: str) -> bytes:
        session = self._session
        if session is None:
            import requests
            session = requests
        self._spacing.wait()
        ua = os.getenv("ISSUER_CSV_USER_AGENT", "Mozilla/5.0 (compatible; portfolio-agent/1.0)")
        try:
            r = session.get(url, timeout=self._timeout, headers={"User-Agent": ua})
        except Exception as exc:
            raise IssuerFileError(f"request failed: {type(exc).__name__}") from None
        if r.status_code != 200:
            raise IssuerFileError(f"HTTP {r.status_code}")
        if len(r.content) > MAX_BYTES:
            raise IssuerFileError("file is too large")
        return r.content

    def fetch(self, ticker: str) -> FundData:
        ticker = ticker.upper()
        src = fdb.get_issuer_source(ticker)                              # enabled rows only
        if src is None:
            return FundData(ticker=ticker, source=self.name, reason_codes=[RC_ISSUER_NOT_CONFIGURED])
        try:
            validate_issuer_url(src["issuer"], src["url"])               # re-checked here: a row edited directly in the database is still held to it
            data = issuer_file_to_fund_data(ticker, parse_issuer_file(self._download(src["url"])))
            return data
        except (IssuerFileError, ValueError):
            return FundData(ticker=ticker, source=self.name, reason_codes=[RC_FUND_FETCH_FAILED])
        except Exception:
            return FundData(ticker=ticker, source=self.name, reason_codes=[RC_FUND_FETCH_FAILED])

    def get_profile(self, ticker): return None
    def get_holdings(self, ticker): return self.fetch(ticker).holdings
    def get_sector_weights(self, ticker): return self.fetch(ticker).sector_weights
    def get_asset_class_weights(self, ticker): return self.fetch(ticker).asset_class_weights
